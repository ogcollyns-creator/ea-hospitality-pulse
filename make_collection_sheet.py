#!/usr/bin/env python3
"""
Build the weekly rate-collection workbook from rates/basket.json + rates/observations.csv.

    python3 make_collection_sheet.py [--week 2026-W39] [--out FILE.xlsx]

WHY THE SHEET LOOKS THE WAY IT DOES. The first live run (2026-W38) produced 75 filled
rows of which 67 were the TWO-NIGHT STAY TOTAL rather than the per-night rate the index
stores — because the sheet said how to make the booking and never said what number to
type. The median week-on-week "move" came out at exactly +100%. Two further faults came
from the same gap: a three-bedroom villa priced instead of the lowest double, and a suite
compared against a prior for a standard room.

So this version removes the judgement from the collector:
  * you type the TOTAL the site shows; the sheet divides by LOS itself
  * ROOM TYPE BOOKED is captured, so a room-type switch is visible instead of silent
  * a ⚠ CHECK flag fires on any per-night move over 30%
  * the convention ("lowest available double") is stated on the entry row, not buried
"""
import argparse, csv, datetime, json, os, collections
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
BASKET = os.path.join(HERE, "rates", "basket.json")
OBS = os.path.join(HERE, "rates", "observations.csv")

F = "Arial"
H1   = Font(name=F, size=14, bold=True, color="FFFFFF")
HDR  = Font(name=F, size=9,  bold=True, color="FFFFFF")
BOLD = Font(name=F, size=10, bold=True)
BODY = Font(name=F, size=10)
MUT  = Font(name=F, size=9,  color="6B6B6B")
INP  = Font(name=F, size=10, color="0000FF")
AUTO = Font(name=F, size=10, color="008000")
WARN = Font(name=F, size=10, bold=True, color="C00000")
TEAL = PatternFill("solid", fgColor="0F6D63")
AMB  = PatternFill("solid", fgColor="C8892F")
GRN  = PatternFill("solid", fgColor="E8F3EC")
SAND = PatternFill("solid", fgColor="F3EFE7")
YEL  = PatternFill("solid", fgColor="FFFF00")
DUE  = PatternFill("solid", fgColor="FFF4CC")
thin = Side(style="thin", color="D5D5D5")
BOX  = Border(left=thin, right=thin, top=thin, bottom=thin)
DATEFMT = "yyyy-mm-dd"

COLS = [("Due this wk",11),("Market",21),("Property",32),("Grp",5),("Last seen",10),("Wks ago",8),
        ("Last US$/night",13),("Prior basis",10),("Prior type",11),("Prior room type",20),
        ("TOTAL FOR STAY US$\n(what the site shows)",18),("PER NIGHT US$\n(auto)",13),
        ("ROOM TYPE BOOKED",22),("NEW BASIS",11),("RATE TYPE",12),
        ("CHECK-IN",12),("CHECK-OUT",12),
        ("SOURCE (url)",28),("NOTE",20),("Changed?",10),("⚠",7),
        ("CSV row to paste into observations.csv",54)]
(C_DUE,C_MKT,C_PROP,C_GRP,C_SEEN,C_AGO,C_LAST,C_PB,C_PT,C_PR,
 C_TOTAL,C_NIGHT,C_ROOM,C_BASIS,C_TYPE,C_IN,C_OUT,C_SRC,C_NOTE,C_CHG,C_FLAG,C_CSV) = range(1,23)
LASTCOL = get_column_letter(C_CSV)

BASIS = [("RO","Room only"),("BB","Bed & breakfast"),("HB","Half board"),("FB","Full board"),
         ("FB+","Full board plus (meals + some activities)"),("AI","All inclusive"),
         ("FI","Fully inclusive (meals, drinks, activities — typical safari basis)")]
RTYPE = [("international","Published international / rack rate"),
         ("resident","East African or national resident rate"),
         ("promotional","Time-limited offer or package rate")]

TIERS = [("City","city",1,"Weekly — every property, every week. Only 30% of city readings were flat; Nairobi just 11%."),
         ("Beach","beach",2,"Fortnightly — half the basket each week. 49% of beach readings were flat."),
         ("Bush","bush",3,"Three-week rotation — a third each week. 82% of bush readings were flat "
                          "(Volcanoes 88%, Mara & Serengeti 81%, Bwindi 78%).")]


def iso(d):
    y, w, _ = d.isocalendar()
    return f"{y}-W{w:02d}"


def history():
    """property -> most recent direct observation."""
    last = {}
    if not os.path.exists(OBS):
        return last
    with open(OBS, encoding="utf-8") as f:
        for r in csv.DictReader(f):
            if r.get("channel", "direct") != "direct":
                continue
            k = (r["market"], r["property"])
            if k not in last or r["observed_date"] > last[k]["observed_date"]:
                last[k] = r
    return last


def room_from_note(note):
    """The generator writes 'Room: X. ...' into the note, so it can be read back."""
    if not note:
        return ""
    if note.startswith("Room: "):
        return note[6:].split(". ")[0][:40]
    return ""


def build(week_tag, obs_date, out_path):
    b = json.load(open(BASKET, encoding="utf-8"))
    last = history()
    wb = Workbook(); wb.remove(wb.active)

    # ---------------- Lists ----------------
    ls = wb.create_sheet("Lists"); ls.sheet_view.showGridLines = False
    for c, t in (("A1","BASIS"),("B1","What it includes"),("D1","RATE TYPE"),("E1","What it means")):
        ls[c] = t; ls[c].font = HDR; ls[c].fill = TEAL
    for i,(k,v) in enumerate(BASIS,2):
        ls.cell(i,1,k).font = BODY; ls.cell(i,2,v).font = MUT
    for i,(k,v) in enumerate(RTYPE,2):
        ls.cell(i,4,k).font = BODY; ls.cell(i,5,v).font = MUT
    for col,w in (("A",10),("B",50),("C",3),("D",14),("E",42)):
        ls.column_dimensions[col].width = w
    ls["A11"] = ("Basis is mandatory. The index must know whether what the price includes has changed — "
                 "a room-only rate and a fully-inclusive rate are not the same product.")
    ls["A11"].font = MUT; ls.merge_cells("A11:B11")

    # ---------------- Control ----------------
    ct = wb.create_sheet("Control", 0); ct.sheet_view.showGridLines = False
    ct.merge_cells("A1:F1"); ct["A1"] = "EA PULSE RATE INDEX — WEEKLY COLLECTION SHEET"
    ct["A1"].font = H1; ct["A1"].fill = TEAL
    ct["A1"].alignment = Alignment(vertical="center", indent=1); ct.row_dimensions[1].height = 28
    for col,w in zip("ABCDEF",(30,16,34,14,14,42)):
        ct.column_dimensions[col].width = w

    def lab(r, t, note="", bold=True):
        ct.cell(r,1,t).font = BOLD if bold else BODY
        if note: ct.cell(r,3,note).font = MUT

    lab(2,"ISO week you are collecting"); ct["B2"] = week_tag
    ct["B2"].font = INP; ct["B2"].fill = YEL; ct["C2"] = "edit (e.g. 2026-W39)"; ct["C2"].font = MUT
    lab(3,"Week counter (auto)", bold=False)
    ct["B3"] = "=VALUE(LEFT(B2,4))*53+VALUE(MID(B2,7,2))"; ct["B3"].font = BODY
    lab(4,"OBSERVATION DATE"); ct["B4"] = obs_date
    ct["B4"].font = INP; ct["B4"].fill = YEL; ct["B4"].number_format = DATEFMT
    ct["C4"] = "change this and the stay dates follow"; ct["C4"].font = MUT
    lab(5,"Lead time (days ahead)"); ct["B5"] = 30; ct["B5"].font = INP; ct["B5"].fill = YEL
    ct["C5"] = "booking convention"; ct["C5"].font = MUT
    lab(6,"Length of stay (nights)"); ct["B6"] = 2; ct["B6"].font = INP; ct["B6"].fill = YEL
    ct["C6"] = "also the divisor for the per-night rate"; ct["C6"].font = MUT
    lab(7,"CHECK-IN DATE (auto)"); ct["B7"] = "=B4+B5"
    ct["B7"].font = AUTO; ct["B7"].fill = GRN; ct["B7"].number_format = DATEFMT
    ct["C7"] = "observation date + lead time"; ct["C7"].font = MUT
    lab(8,"CHECK-OUT DATE (auto)"); ct["B8"] = "=B7+B6"
    ct["B8"].font = AUTO; ct["B8"].fill = GRN; ct["B8"].number_format = DATEFMT
    ct["C8"] = "check-in + nights"; ct["C8"].font = MUT
    for r in range(2,9):
        for c in range(1,4): ct.cell(r,c).border = BOX

    ct["A10"] = "WHAT TO PRICE"; ct["A10"].font = Font(name=F, size=12, bold=True, color="0F6D63")
    rules = ["The LOWEST AVAILABLE DOUBLE room, 2 adults, for the check-in and check-out dates above.",
             "NOT a suite, villa, cottage or upgrade — even if it is the only thing showing. If no standard",
             "double is available, leave the row blank and say so in NOTE. A blank row costs nothing; a villa",
             "priced against last week's standard room invents a rate rise that never happened.",
             "Type the TOTAL the site shows for the whole stay in column K. The sheet works out the per-night",
             "rate itself — you never divide anything.",
             "Record the room name you actually booked in ROOM TYPE BOOKED, so a switch is visible later."]
    for n,t in enumerate(rules):
        c = ct.cell(11+n,1,t); c.font = BODY if n < 1 else MUT
        ct.merge_cells(start_row=11+n, start_column=1, end_row=11+n, end_column=6)

    ct["A19"] = "THE THREE TIERS"; ct["A19"].font = Font(name=F, size=12, bold=True, color="0F6D63")
    for i,h in enumerate(["Tier","Properties","Cadence","Per week","% flat","Why"],1):
        c = ct.cell(20,i,h); c.font = HDR; c.fill = TEAL; c.border = BOX
        c.alignment = Alignment(horizontal="center", wrap_text=True)
    ct.row_dimensions[20].height = 26
    counts = collections.Counter()
    for mk,m in b["markets"].items():
        counts[m["segment"]] += len(m["properties"])
    meta = [("City","city",1,"Weekly","30%","Rates reprice continuously. Nairobi moved in 89% of readings."),
            ("Beach","beach",2,"Fortnightly","49%","Moves on shoulder/peak turns, not weekly."),
            ("Bush","bush",3,"3-week rotation","82%","Seasonal rate sheets — a weekly re-read confirmed no change 4 times in 5.")]
    for n,(t,seg,ev,cad,flat,why) in enumerate(meta):
        r = 21+n
        ct.cell(r,1,t).font = BOLD; ct.cell(r,2,counts[seg]); ct.cell(r,3,cad)
        ct.cell(r,4,f"=ROUND(B{r}/{ev},0)"); ct.cell(r,5,flat); ct.cell(r,6,why)
        for i in range(1,7):
            c = ct.cell(r,i); c.border = BOX
            if i != 1: c.font = BODY
            if i in (2,3,4,5): c.alignment = Alignment(horizontal="center")
        ct.cell(r,6).font = MUT; ct.cell(r,6).alignment = Alignment(wrap_text=True, vertical="top")
        ct.row_dimensions[r].height = 28
    ct.cell(24,1,"TOTAL").font = BOLD
    ct.cell(24,2,"=SUM(B21:B23)").font = BOLD
    ct.cell(24,4,"=SUM(D21:D23)").font = BOLD
    ct.cell(24,6,"vs 136/week on a full weekly sweep").font = MUT
    for i in range(1,7):
        ct.cell(24,i).border = BOX; ct.cell(24,i).fill = SAND
    for i in (2,4): ct.cell(24,i).alignment = Alignment(horizontal="center")

    ct["A26"] = "HOW TO USE"; ct["A26"].font = Font(name=F, size=12, bold=True, color="0F6D63")
    steps = ["1. Set the ISO week (B2) and the OBSERVATION DATE (B4). Check-in and check-out compute themselves.",
             "2. Open City, Beach and Bush. Filter column A ('Due this wk') to YES.",
             "3. Fill the BLUE columns. K is the stay total from the site — the sheet derives the per-night rate.",
             "4. A ⚠ CHECK appears if the per-night rate moved more than 30%. Confirm the room type before",
             "    accepting it; that flag is usually a room-type switch, not a price move.",
             "5. Copy the CSV column into rates/observations.csv, then run:  python3 build_rate_index.py",
             "",
             "Why the cadence is safe: build_rate_index.py compares every property only with ITSELF and matches",
             "back up to LOOKBACK_WEEKS = 6, so a 2- or 3-week gap chains exactly like weekly collection."]
    for n,t in enumerate(steps):
        c = ct.cell(27+n,1,t); c.font = MUT if n >= 6 else BODY
        ct.merge_cells(start_row=27+n, start_column=1, end_row=27+n, end_column=6)

    # ---------------- tier tabs ----------------
    made = []
    for title, seg, every, blurb in TIERS:
        ws = wb.create_sheet(title); ws.sheet_view.showGridLines = False
        ws.merge_cells(f"A1:{LASTCOL}1"); c = ws["A1"]
        c.value = f"{title.upper()} — cadence: every {every} week{'s' if every>1 else ''}"
        c.font = H1; c.fill = TEAL
        c.alignment = Alignment(vertical="center", indent=1); ws.row_dimensions[1].height = 26
        ws.merge_cells(f"A2:{LASTCOL}2"); ws["A2"] = blurb
        ws["A2"].font = MUT; ws["A2"].alignment = Alignment(indent=1)
        ws.merge_cells(f"A3:{LASTCOL}3")
        ws["A3"] = ("Price the LOWEST AVAILABLE DOUBLE, 2 adults, for the CHECK-IN / CHECK-OUT dates. "
                    "In column K type the TOTAL the site shows for the stay — the sheet works out the per-night "
                    "rate (column L). BLUE = you type it. GREEN = automatic. Filter column A to YES.")
        ws["A3"].font = Font(name=F, size=9, italic=True, color="B35C00")
        ws["A3"].alignment = Alignment(indent=1)

        hr = 5
        for i,(h,w) in enumerate(COLS,1):
            cell = ws.cell(hr,i,h); cell.font = HDR
            cell.fill = AMB if i in (C_NIGHT,C_IN,C_OUT) else TEAL
            cell.border = BOX
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.row_dimensions[hr].height = 34

        data = []
        for mk,m in b["markets"].items():
            if m["segment"] != seg: continue
            grp = ({p:g for g,ps in m.get("rotation",{}).items() for p in ps} if every > 1 else {})
            for p in m["properties"]:
                h = last.get((mk,p))
                data.append(dict(mk=mk, label=m["label"], prop=p,
                                 group=int(grp.get(p,0)) if every > 1 else 0,
                                 week=iso(datetime.date.fromisoformat(h["observed_date"])) if h else "",
                                 rate=float(h["rate_usd"]) if h else "",
                                 basis=h["basis"] if h else "", rtype=h["rate_type"] if h else "",
                                 room=room_from_note(h["note"]) if h else ""))
        data.sort(key=lambda d:(d["label"], d["group"], d["prop"]))

        r0 = hr+1
        for n,d in enumerate(data):
            r = r0+n
            ws.cell(r,C_DUE,"YES" if every == 1 else f'=IF(MOD(Control!$B$3,{every})=D{r},"YES","—")')
            ws.cell(r,C_MKT,d["label"]); ws.cell(r,C_PROP,d["prop"])
            ws.cell(r,C_GRP,d["group"] if every > 1 else "")
            ws.cell(r,C_SEEN,d["week"] or "never")
            ws.cell(r,C_AGO,f'=IF(E{r}="never","",Control!$B$3-(VALUE(LEFT(E{r},4))*53+VALUE(MID(E{r},7,2))))')
            ws.cell(r,C_LAST,d["rate"]); ws.cell(r,C_PB,d["basis"])
            ws.cell(r,C_PT,d["rtype"]); ws.cell(r,C_PR,d["room"])
            for col in (C_TOTAL,C_ROOM,C_BASIS,C_TYPE,C_SRC,C_NOTE):
                cc = ws.cell(r,col); cc.font = INP; cc.fill = YEL
            ws.cell(r,C_TOTAL).number_format = "#,##0.00"
            pn = ws.cell(r,C_NIGHT,f'=IF(K{r}="","",ROUND(K{r}/Control!$B$6,2))')
            pn.font = AUTO; pn.fill = GRN; pn.number_format = "#,##0.00"
            for col,src in ((C_IN,"$B$7"),(C_OUT,"$B$8")):
                cc = ws.cell(r,col,f'=IF($K{r}="","",Control!{src})')
                cc.font = AUTO; cc.fill = GRN; cc.number_format = DATEFMT
                cc.alignment = Alignment(horizontal="center")
            ws.cell(r,C_CHG,f'=IF(L{r}="","",IF(G{r}="","new",'
                            f'IF(L{r}=G{r},"same",TEXT((L{r}-G{r})/G{r},"+0.0%;-0.0%"))))')
            ws.cell(r,C_FLAG,f'=IF(OR(L{r}="",G{r}=""),"",IF(ABS(L{r}/G{r}-1)>0.3,"CHECK",""))')
            ws.cell(r,C_CSV,
                '=IF(L{r}="","",_xlfn.TEXTJOIN(",",FALSE,TEXT(Control!$B$4,"yyyy-mm-dd"),"{mk}",'
                '""""&C{r}&"""",L{r},TEXT(P{r},"yyyy-mm-dd"),Control!$B$6,""""&R{r}&"""",'
                '""""&"Room: "&M{r}&". Stay total USD "&K{r}&" over "&Control!$B$6&" nights. "&S{r}&"""",'
                'N{r},O{r},"direct"))'.format(r=r, mk=d["mk"]))
            for col in range(1,C_CSV+1):
                cell = ws.cell(r,col); cell.border = BOX
                if col not in (C_TOTAL,C_ROOM,C_BASIS,C_TYPE,C_SRC,C_NOTE,C_NIGHT,C_IN,C_OUT):
                    cell.font = BODY
                if col in (C_DUE,C_GRP,C_AGO,C_CHG,C_FLAG):
                    cell.alignment = Alignment(horizontal="center")
                if col == C_LAST: cell.number_format = "#,##0.00"
                if col == C_CSV: cell.font = Font(name="Consolas", size=8, color="6B6B6B")
            ws.cell(r,C_DUE).font = BOLD
            ws.cell(r,C_FLAG).font = WARN

        last_row = r0+len(data)-1
        ws.conditional_formatting.add(f"A{r0}:{LASTCOL}{last_row}",
            FormulaRule(formula=[f'$A{r0}="YES"'], fill=DUE))
        fl = get_column_letter(C_FLAG)
        ws.conditional_formatting.add(f"{fl}{r0}:{fl}{last_row}",
            FormulaRule(formula=[f'${fl}{r0}="CHECK"'], fill=PatternFill("solid", fgColor="FFC7CE")))

        dvb = DataValidation(type="list", formula1="Lists!$A$2:$A$8", allow_blank=True, showDropDown=False,
            showErrorMessage=True, errorTitle="Pick a basis",
            error="Choose RO, BB, HB, FB, FB+, AI or FI. See the Lists tab.",
            showInputMessage=True, promptTitle="Basis", prompt="RO / BB / HB / FB / FB+ / AI / FI")
        ws.add_data_validation(dvb)
        dvb.add(f"{get_column_letter(C_BASIS)}{r0}:{get_column_letter(C_BASIS)}{last_row}")
        dvt = DataValidation(type="list", formula1="Lists!$D$2:$D$4", allow_blank=True, showDropDown=False,
            showErrorMessage=True, errorTitle="Pick a rate type",
            error="Choose international, resident or promotional.",
            showInputMessage=True, promptTitle="Rate type", prompt="international / resident / promotional")
        ws.add_data_validation(dvt)
        dvt.add(f"{get_column_letter(C_TYPE)}{r0}:{get_column_letter(C_TYPE)}{last_row}")

        ws.freeze_panes = f"D{r0}"
        ws.auto_filter.ref = f"A{hr}:{LASTCOL}{last_row}"
        tot = last_row+2
        ws.cell(tot,C_PROP,"Due this week:").font = BOLD
        ws.cell(tot,C_GRP,f'=COUNTIF(A{r0}:A{last_row},"YES")').font = BOLD
        ws.cell(tot,C_SEEN,"Filled (of due):").font = BOLD
        ws.cell(tot,C_AGO,f'=COUNTIFS(A{r0}:A{last_row},"YES",K{r0}:K{last_row},"<>")').font = BOLD
        ws.cell(tot,C_LAST,"Still to do:").font = BOLD
        ws.cell(tot,C_PB,f'=COUNTIFS(A{r0}:A{last_row},"YES",K{r0}:K{last_row},"")').font = BOLD
        ws.cell(tot,C_ROOM,"Flagged ⚠:").font = BOLD
        ws.cell(tot,C_BASIS,f'=COUNTIF(U{r0}:U{last_row},"CHECK")').font = WARN
        made.append((title,len(data)))

    wb.move_sheet("Lists", offset=3)
    wb.save(out_path)
    return made


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--week", help="ISO week, e.g. 2026-W39 (default: current)")
    ap.add_argument("--date", help="observation date YYYY-MM-DD (default: today)")
    ap.add_argument("--out", default="EA-Pulse-Rate-Index-Collection.xlsx")
    a = ap.parse_args()
    d = datetime.date.fromisoformat(a.date) if a.date else datetime.date.today()
    tag = a.week or iso(d)
    made = build(tag, d, a.out)
    print(f"{a.out} — week {tag}, observation date {d}")
    for t,n in made:
        print(f"  {t}: {n} properties")
